"""Create a distribution from a spreadsheet.

One row per category entry, with a header row. Columns, by header name (case
and spacing don't matter; "Min TU", "min_tossups" and "Minimum Tossups" all
work):

    Category | Subcategory | Min Tossups | Max Tossups | Min Bonuses | Max Bonuses

Subcategory may be blank. A blank max takes the min; a blank min takes the
max; both blank is 0. Counts are per packet. Accepts .xlsx, .csv and .tsv.
"""

import csv
import io
import re

from django.utils import timezone

from .models import Distribution, DistributionEntry


class DistributionImportError(Exception):
    pass


COLUMNS = [
    ('category', 'Category'),
    ('subcategory', 'Subcategory'),
    ('min_tossups', 'Min Tossups'),
    ('max_tossups', 'Max Tossups'),
    ('min_bonuses', 'Min Bonuses'),
    ('max_bonuses', 'Max Bonuses'),
]

# The template's example rows: a plausible 20/20 packet in miniature.
TEMPLATE_ROWS = [
    ('Literature', 'American', 1, 1, 1, 1),
    ('Literature', 'British', 1, 1, 1, 1),
    ('Literature', 'World', 1, 2, 1, 2),
    ('History', 'American', 1, 1, 1, 1),
    ('History', 'European', 1, 1, 1, 1),
    ('History', 'World', 1, 2, 1, 2),
    ('Science', 'Biology', 1, 1, 1, 1),
    ('Science', 'Chemistry', 1, 1, 1, 1),
    ('Science', 'Physics', 1, 1, 1, 1),
    ('Science', 'Other', 0, 1, 0, 1),
    ('Fine Arts', 'Visual', 1, 1, 1, 1),
    ('Fine Arts', 'Audio', 1, 1, 1, 1),
    ('Fine Arts', 'Other', 0, 1, 0, 1),
    ('Religion', '', 1, 1, 1, 1),
    ('Mythology', '', 1, 1, 1, 1),
    ('Philosophy', '', 1, 1, 1, 1),
    ('Social Science', '', 1, 1, 1, 1),
    ('Geography', '', 1, 1, 1, 1),
    ('Current Events', '', 1, 1, 1, 1),
    ('Trash', '', 1, 1, 1, 1),
]


def _norm_header(text):
    """'Min Tossups', 'min_tossups', 'Minimum TU' -> 'min_tossups'."""
    t = re.sub(r'[^a-z]+', ' ', (text or '').strip().lower()).strip()
    t = t.replace('minimum', 'min').replace('maximum', 'max')
    t = re.sub(r'\btossup\b|\btu\b|\btus\b', 'tossups', t)
    t = re.sub(r'\bbonus\b|\bbs\b|\bb\b', 'bonuses', t)
    t = t.replace(' ', '_')
    aliases = {
        'cat': 'category', 'subcat': 'subcategory', 'sub_category': 'subcategory',
        'sub': 'subcategory', 'min_tossups': 'min_tossups', 'max_tossups': 'max_tossups',
        'min_bonuses': 'min_bonuses', 'max_bonuses': 'max_bonuses',
        'tossups': 'max_tossups', 'bonuses': 'max_bonuses',
    }
    return aliases.get(t, t)


def _rows_from_upload(name, data):
    """The sheet as a list of rows of strings, from .xlsx or delimited text."""
    lower = (name or '').lower()
    if lower.endswith('.xlsx') or lower.endswith('.xlsm'):
        try:
            import openpyxl
        except ImportError:
            raise DistributionImportError('Excel files need the openpyxl package; upload a CSV instead.')
        try:
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception:
            raise DistributionImportError('That does not look like an Excel (.xlsx) file.')
        ws = wb.worksheets[0]
        return [['' if v is None else str(v).strip() for v in row] for row in ws.iter_rows(values_only=True)]
    if lower.endswith('.xls'):
        raise DistributionImportError('Old .xls files are not supported; save as .xlsx or .csv.')
    try:
        text = data.decode('utf-8-sig')
    except UnicodeDecodeError:
        text = data.decode('latin-1')
    delimiter = '\t' if lower.endswith('.tsv') or lower.endswith('.txt') or ('\t' in text and ',' not in text) else ','
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _int_or_none(value, what, line):
    v = (value or '').strip()
    if v == '':
        return None
    try:
        n = float(v)
    except ValueError:
        raise DistributionImportError('Line {0}: {1} should be a number, not "{2}"'.format(line, what, v))
    if n < 0 or n != int(n):
        raise DistributionImportError('Line {0}: {1} should be a whole number, not "{2}"'.format(line, what, v))
    return int(n)


def parse_distribution_sheet(name, data):
    """-> list of entry dicts (category, subcategory, min/max tossups/bonuses).
    Raises DistributionImportError with a line-numbered message on bad input."""
    rows = _rows_from_upload(name, data)
    rows = [r for r in rows if any((c or '').strip() for c in r)]
    if not rows:
        raise DistributionImportError('The sheet is empty.')
    header = [_norm_header(c) for c in rows[0]]
    if 'category' not in header:
        raise DistributionImportError(
            'The first row must be a header naming the columns -- at least "Category". '
            'Download the template to see the layout.')
    col = {key: header.index(key) for key, _ in COLUMNS if key in header}
    if not any(k in col for k in ('min_tossups', 'max_tossups', 'min_bonuses', 'max_bonuses')):
        raise DistributionImportError('No count columns found (Min/Max Tossups, Min/Max Bonuses).')

    def cell(row, key):
        i = col.get(key)
        return row[i] if i is not None and i < len(row) else ''

    entries, seen = [], set()
    for n, row in enumerate(rows[1:], start=2):
        category = cell(row, 'category').strip()
        subcategory = cell(row, 'subcategory').strip()
        if not category:
            raise DistributionImportError('Line {0}: the category is blank'.format(n))
        key = (category.lower(), subcategory.lower())
        if key in seen:
            raise DistributionImportError('Line {0}: "{1}{2}" appears more than once'.format(
                n, category, ' - ' + subcategory if subcategory else ''))
        seen.add(key)
        min_tu = _int_or_none(cell(row, 'min_tossups'), 'min tossups', n)
        max_tu = _int_or_none(cell(row, 'max_tossups'), 'max tossups', n)
        min_bs = _int_or_none(cell(row, 'min_bonuses'), 'min bonuses', n)
        max_bs = _int_or_none(cell(row, 'max_bonuses'), 'max bonuses', n)
        # A blank side takes the other; both blank is 0.
        if max_tu is None:
            max_tu = min_tu if min_tu is not None else 0
        if min_tu is None:
            min_tu = max_tu
        if max_bs is None:
            max_bs = min_bs if min_bs is not None else 0
        if min_bs is None:
            min_bs = max_bs
        if min_tu > max_tu:
            raise DistributionImportError('Line {0}: min tossups ({1}) is more than max ({2})'.format(n, min_tu, max_tu))
        if min_bs > max_bs:
            raise DistributionImportError('Line {0}: min bonuses ({1}) is more than max ({2})'.format(n, min_bs, max_bs))
        entries.append({'category': category, 'subcategory': subcategory,
                        'min_tossups': min_tu, 'max_tossups': max_tu,
                        'min_bonuses': min_bs, 'max_bonuses': max_bs})
    if not entries:
        raise DistributionImportError('The sheet has a header but no entries under it.')
    return entries


def create_distribution_from_sheet(name, data, dist_name, owner, public=False):
    """Parse the sheet and create the distribution with its entries.
    Returns the Distribution."""
    entries = parse_distribution_sheet(name, data)
    dist = Distribution(name=dist_name.strip()[:100], public=public,
                        created_by=owner, created_date=timezone.now())
    dist.acf_tossup_per_period_count = sum(e['max_tossups'] for e in entries)
    dist.acf_bonus_per_period_count = sum(e['max_bonuses'] for e in entries)
    dist.save()
    for e in entries:
        DistributionEntry.objects.create(distribution=dist, **e)
    return dist


def template_rows():
    """Header plus example rows, for the downloadable templates."""
    return [[label for _, label in COLUMNS]] + [list(r) for r in TEMPLATE_ROWS]


def template_csv():
    out = io.StringIO()
    w = csv.writer(out)
    for row in template_rows():
        w.writerow(row)
    return out.getvalue()


def template_xlsx():
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Distribution'
    for row in template_rows():
        ws.append(row)
    for c in ws[1]:
        c.font = Font(bold=True)
    for col, width in zip('ABCDEF', (18, 16, 13, 13, 13, 13)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = 'A2'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
